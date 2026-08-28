"""Preview-first master-data imports.

Uploads are parsed into normalized rows and errors before any database mutation.
Callers commit only when the complete file is valid, making each import atomic.
"""

from __future__ import annotations

import csv
import io
import zipfile
from dataclasses import dataclass
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.domain.identity import InvalidPhoneNumber, normalize_indian_whatsapp_number
from app.models import CompetitionPrice, Employee, EmploymentType, PriceType, Product, ProductOwnership

MAX_IMPORT_ROWS = 20_000
MAX_IMPORT_COLUMNS = 64
MAX_XLSX_UNCOMPRESSED_BYTES = 50_000_000


def _headers(values: list[Any]) -> list[str]:
    headers = [str(value or "").strip() for value in values]
    if not headers or any(not header for header in headers):
        raise ValueError("Every column must have a non-empty header")
    folded = [header.casefold() for header in headers]
    if len(folded) != len(set(folded)):
        raise ValueError("Column headers must be unique")
    if len(headers) > MAX_IMPORT_COLUMNS:
        raise ValueError(f"Master files may contain at most {MAX_IMPORT_COLUMNS} columns")
    return headers


@dataclass(frozen=True)
class ImportError:
    row: int
    field: str
    message: str


@dataclass(frozen=True)
class ImportPreview:
    rows: list[dict[str, Any]]
    errors: list[ImportError]

    def as_dict(self) -> dict[str, Any]:
        return {
            "valid": not self.errors,
            "row_count": len(self.rows),
            "rows": self.rows,
            "errors": [error.__dict__ for error in self.errors],
        }


def read_tabular_upload(filename: str, content: bytes) -> list[dict[str, str]]:
    if filename.casefold().endswith(".xlsx"):
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(item.file_size for item in archive.infolist()) > MAX_XLSX_UNCOMPRESSED_BYTES:
                    raise ValueError("XLSX content exceeds the 50 MB decompressed limit")
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            try:
                sheet = workbook.active
                # Cached dimensions are optional in XLSX. Use them only as a
                # fast rejection path; forcing their calculation can scan an
                # attacker-controlled sparse sheet before our limits apply.
                if (sheet.max_column is not None and sheet.max_column > MAX_IMPORT_COLUMNS) or (
                    sheet.max_row is not None and sheet.max_row > MAX_IMPORT_ROWS + 1
                ):
                    raise ValueError(
                        f"Master files may contain at most {MAX_IMPORT_ROWS} rows and {MAX_IMPORT_COLUMNS} columns"
                    )
                iterator = sheet.iter_rows(values_only=True)
                first_row = next(iterator, None)
                if first_row is None:
                    return []
                headers = _headers(list(first_row))
                rows: list[dict[str, str]] = []
                for sheet_row_number, row in enumerate(iterator, start=2):
                    # Count yielded worksheet rows, including blanks, so a
                    # dimensionless sparse coordinate remains CPU-bounded.
                    if sheet_row_number > MAX_IMPORT_ROWS + 1:
                        raise ValueError(f"Master files may contain at most {MAX_IMPORT_ROWS} data rows")
                    if len(row) > MAX_IMPORT_COLUMNS:
                        raise ValueError(f"Master files may contain at most {MAX_IMPORT_COLUMNS} columns")
                    if any(value not in (None, "") for value in row[len(headers) :]):
                        raise ValueError("A data row contains more values than the header")
                    if not any(value not in (None, "") for value in row):
                        continue
                    rows.append(
                        {
                            header: str((row[index] if index < len(row) else "") or "").strip()
                            for index, header in enumerate(headers)
                        }
                    )
                return rows
            finally:
                workbook.close()
        except (zipfile.BadZipFile, KeyError, OSError) as exc:
            raise ValueError("The XLSX file is malformed") from exc
    text = content.decode("utf-8-sig")
    try:
        reader = csv.DictReader(io.StringIO(text))
        headers = _headers(list(reader.fieldnames or []))
        rows = []
        for row in reader:
            if len(rows) >= MAX_IMPORT_ROWS:
                raise ValueError(f"Master files may contain at most {MAX_IMPORT_ROWS} data rows")
            if None in row:
                raise ValueError("A data row contains more values than the header")
            rows.append({header: (row.get(header) or "").strip() for header in headers})
        return rows
    except csv.Error as exc:
        raise ValueError("The CSV file is malformed") from exc


def _value(row: dict[str, str], *names: str) -> str:
    folded = {key.casefold().strip(): value for key, value in row.items()}
    for name in names:
        if name.casefold() in folded:
            return folded[name.casefold()].strip()
    return ""


def preview_employees(rows: list[dict[str, str]]) -> ImportPreview:
    normalized: list[dict[str, Any]] = []
    errors: list[ImportError] = []
    if not rows:
        return ImportPreview([], [ImportError(1, "file", "The employee master contains no data rows")])
    batch_numbers: dict[str, int] = {}
    batch_codes: dict[str, int] = {}
    for row_number, row in enumerate(rows, start=2):
        code = _value(row, "employee_code", "employee code")
        raw_number = _value(row, "whatsapp_number", "whatsapp phone number", "whatsapp number")
        employment = _value(row, "employment_type", "employment type", "full time or contractual").casefold().replace(" ", "_")
        state = _value(row, "state")
        if not code:
            errors.append(ImportError(row_number, "employee_code", "Employee code is required"))
        if not state:
            errors.append(ImportError(row_number, "state", "State is required"))
        if employment in {"fulltime", "full_time", "full-time"}:
            employment = EmploymentType.FULL_TIME.value
        elif employment in {"contract", "contractual"}:
            employment = EmploymentType.CONTRACTUAL.value
        else:
            errors.append(ImportError(row_number, "employment_type", "Use full_time or contractual"))
        try:
            whatsapp_number = normalize_indian_whatsapp_number(raw_number)
        except InvalidPhoneNumber as exc:
            errors.append(ImportError(row_number, "whatsapp_number", str(exc)))
            whatsapp_number = ""
        if code in batch_codes:
            errors.append(ImportError(row_number, "employee_code", f"Duplicate of row {batch_codes[code]}"))
        elif code:
            batch_codes[code] = row_number
        if whatsapp_number in batch_numbers:
            errors.append(ImportError(row_number, "whatsapp_number", f"Duplicate of row {batch_numbers[whatsapp_number]}"))
        elif whatsapp_number:
            batch_numbers[whatsapp_number] = row_number
        normalized.append(
            {
                "employee_code": code,
                "employment_type": employment,
                "designation": _value(row, "designation") or None,
                "territory_code": _value(row, "territory_code", "territory identification", "territory") or None,
                "state": state,
                "email": _value(row, "email", "email id") or None,
                "whatsapp_number": whatsapp_number,
                "active": _value(row, "active").casefold() not in {"false", "no", "0", "inactive"},
            }
        )
    return ImportPreview(normalized, errors)


def commit_employees(session: Session, company_id: str, preview: ImportPreview) -> int:
    if preview.errors:
        raise ValueError("Cannot commit an invalid employee import")
    changed = 0
    for data in preview.rows:
        employee = session.scalar(
            select(Employee).where(Employee.company_id == company_id, Employee.employee_code == data["employee_code"])
        )
        conflicting = session.scalar(
            select(Employee).where(
                Employee.company_id == company_id,
                Employee.whatsapp_number == data["whatsapp_number"],
                Employee.employee_code != data["employee_code"],
                Employee.active.is_(True),
            )
        )
        if conflicting:
            raise ValueError(f"WhatsApp number already belongs to active employee {conflicting.employee_code}")
        if employee is None:
            employee = Employee(company_id=company_id, **data)
            session.add(employee)
        else:
            for field, value in data.items():
                setattr(employee, field, value)
        changed += 1
    return changed


def preview_products(rows: list[dict[str, str]], ownership: ProductOwnership) -> ImportPreview:
    normalized: list[dict[str, Any]] = []
    errors: list[ImportError] = []
    seen: dict[str, int] = {}
    for row_number, row in enumerate(rows, start=2):
        brand = _value(row, "brand", "brand name")
        category = _value(row, "category", "product category")
        if not brand:
            errors.append(ImportError(row_number, "brand", "Brand is required"))
        if not category:
            errors.append(ImportError(row_number, "category", "Category is required"))
        key = brand.casefold()
        if key in seen:
            errors.append(ImportError(row_number, "brand", f"Duplicate of row {seen[key]}"))
        elif brand:
            seen[key] = row_number
        normalized.append(
            {
                "ownership": ownership.value,
                "manufacturer": _value(row, "manufacturer", "company") or None,
                "brand": brand,
                "category": category,
                "active_ingredient": _value(row, "active_ingredient", "active ingredient") or None,
                "formulation": _value(row, "formulation") or None,
                "concentration": _value(row, "concentration") or None,
                "metadata_json": {
                    "pack_sizes": _value(row, "pack_sizes", "pack sizes"),
                    "crops": _value(row, "crops"),
                    "targets": _value(row, "targets", "target pests diseases"),
                    "aliases": _value(row, "aliases", "regional aliases"),
                },
                "active": True,
            }
        )
    return ImportPreview(normalized, errors)


def commit_products(session: Session, company_id: str, preview: ImportPreview) -> int:
    if preview.errors:
        raise ValueError("Cannot commit an invalid product import")
    for data in preview.rows:
        ownership = ProductOwnership(data["ownership"])
        product = session.scalar(
            select(Product).where(
                Product.company_id == company_id,
                Product.ownership == ownership,
                Product.brand == data["brand"],
            )
        )
        values = {**data, "ownership": ownership}
        if product is None:
            session.add(Product(company_id=company_id, **values))
        else:
            for field, value in values.items():
                setattr(product, field, value)
    return len(preview.rows)


def preview_prices(session: Session, company_id: str, rows: list[dict[str, str]]) -> ImportPreview:
    normalized: list[dict[str, Any]] = []
    errors: list[ImportError] = []
    for row_number, row in enumerate(rows, start=2):
        brand = _value(row, "brand", "brand name")
        product = session.scalar(
            select(Product).where(
                Product.company_id == company_id,
                Product.ownership == ProductOwnership.COMPETITOR,
                Product.brand == brand,
            )
        )
        if not product:
            errors.append(ImportError(row_number, "brand", "Competitor brand is not in the product master"))
        price_type_text = _value(row, "price_type", "price type").casefold().replace(" ", "_")
        aliases = {
            "farmer": PriceType.FARMER,
            "farmer_price": PriceType.FARMER,
            "channel_net_landing": PriceType.CHANNEL_NET_LANDING,
            "net_landing_price": PriceType.CHANNEL_NET_LANDING,
        }
        price_type = aliases.get(price_type_text)
        if not price_type:
            errors.append(ImportError(row_number, "price_type", "Use farmer_price or channel_net_landing"))
        try:
            amount = Decimal(_value(row, "amount", "price"))
            if amount <= 0:
                raise InvalidOperation
        except (InvalidOperation, ValueError):
            errors.append(ImportError(row_number, "amount", "Amount must be greater than zero"))
            amount = Decimal("0")
        state = _value(row, "state")
        pack_size = _value(row, "pack_size", "pack size")
        if not state:
            errors.append(ImportError(row_number, "state", "State is required"))
        if not pack_size:
            errors.append(ImportError(row_number, "pack_size", "Pack size is required"))
        observed_text = _value(row, "observed_at", "effective date", "date")
        try:
            observed_at = datetime.fromisoformat(observed_text)
            if observed_at.tzinfo is None:
                observed_at = observed_at.replace(tzinfo=timezone.utc)
            else:
                observed_at = observed_at.astimezone(timezone.utc)
        except ValueError:
            errors.append(ImportError(row_number, "observed_at", "A valid ISO effective date is required"))
            observed_at = datetime.now(timezone.utc)
        normalized.append(
            {
                "product_id": product.id if product else "",
                "state": state,
                "pack_size": pack_size,
                "price_type": price_type.value if price_type else "",
                "amount": str(amount),
                "currency": _value(row, "currency") or "INR",
                "observed_at": observed_at.isoformat(),
                "source": "initial_upload",
            }
        )
    return ImportPreview(normalized, errors)


def commit_prices(session: Session, company_id: str, preview: ImportPreview) -> int:
    if preview.errors:
        raise ValueError("Cannot commit an invalid price import")
    for data in preview.rows:
        session.add(
            CompetitionPrice(
                company_id=company_id,
                product_id=data["product_id"],
                state=data["state"],
                pack_size=data["pack_size"],
                price_type=PriceType(data["price_type"]),
                amount=Decimal(data["amount"]),
                currency=data["currency"],
                observed_at=datetime.fromisoformat(data["observed_at"]),
                source=data["source"],
                evidence_ids=[],
            )
        )
    return len(preview.rows)
